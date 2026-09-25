from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "Contoh_Perbaikan_Template_Excel_Unlocked.xlsx"

NAVY = "17365D"
BLUE = "D9EAF7"
LIGHT_BLUE = "EAF3F8"
YELLOW = "FFF2CC"
GREEN = "E2F0D9"
RED = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
DARK = "243447"
TEAL = "0F6B78"

thin_gray = Side(style="thin", color="B7C9D6")
border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)


def title(ws, text, subtitle=None):
    ws.merge_cells("A1:J1")
    ws["A1"] = text
    ws["A1"].font = Font(size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    if subtitle:
        ws.merge_cells("A2:J2")
        ws["A2"] = subtitle
        ws["A2"].font = Font(size=10, italic=True, color="5B6573")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = 28


def section(cell, text):
    cell.value = text
    cell.font = Font(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=TEAL)
    cell.alignment = Alignment(vertical="center")


def style_header(row):
    for cell in row:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border


def apply_grid(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def make_start_here(wb):
    ws = wb.active
    ws.title = "START HERE"
    title(ws, "CONTOH PERBAIKAN WORKBOOK", "Contoh ini dibuat ulang dari pengalaman sebelumnya. Semua nama dan data pengukuran bersifat fiktif.")
    ws.sheet_view.showGridLines = False

    ws["A4"] = "Yang diperlihatkan dalam contoh ini"
    section(ws["A4"], ws["A4"].value)
    ws.merge_cells("A5:J6")
    ws["A5"] = (
        "Teknisi mengisi data di satu sheet. Informasi alat diambil dari MASTER DATA, lalu workbook menghitung "
        "hasil dan memperbarui laporan. Reviewer tetap dapat mengikuti angka yang dipakai dalam perhitungan."
    )
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["A5"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    ws["A8"] = "Alur workbook"
    section(ws["A8"], ws["A8"].value)
    flow = [
        ("A10:B12", "1  INPUT LOG\nDiisi oleh teknisi", YELLOW),
        ("D10:E12", "2  MASTER DATA\nReferensi alat", GREEN),
        ("G10:H12", "3  CALCULATION\nRincian hitung", BLUE),
        ("I10:J12", "4  REPORT\nRingkasan hasil", GRAY),
    ]
    for area, text, color in flow:
        ws.merge_cells(area)
        c = ws[area.split(":")[0]]
        c.value = text
        c.font = Font(bold=True, color=DARK)
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=Side(style="medium", color=TEAL), right=Side(style="medium", color=TEAL), top=Side(style="medium", color=TEAL), bottom=Side(style="medium", color=TEAL))
    for arrow_cell in ["C11", "F11"]:
        ws[arrow_cell] = "→"
        ws[arrow_cell].font = Font(size=22, bold=True, color=TEAL)
        ws[arrow_cell].alignment = Alignment(horizontal="center")

    ws["A15"] = "Cara menggunakan"
    section(ws["A15"], ws["A15"].value)
    instructions = [
        "Buka INPUT LOG. Teknisi hanya perlu mengisi sel berwarna kuning.",
        "Pilih Asset ID dan nama teknisi dari daftar yang tersedia.",
        "Isi tiga hasil pengukuran. Workbook akan menghitung rata-rata, selisih, dan status.",
        "Buka REPORT untuk melihat rekap. Hasil kasus acuan ada di VALIDATION.",
        "File demo tidak dikunci agar rumus dapat diperiksa. Proteksi dapat diaktifkan saat masuk tahap UAT.",
    ]
    for idx, text in enumerate(instructions, start=16):
        ws.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=10)
        ws.cell(idx, 1, f"{idx - 15}. {text}")
        ws.cell(idx, 1).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(idx, 1).fill = PatternFill("solid", fgColor=WHITE if idx % 2 == 0 else LIGHT_BLUE)

    ws["A23"] = "Legenda"
    section(ws["A23"], ws["A23"].value)
    legends = [
        ("A24:B25", YELLOW, "Input pengguna"),
        ("D24:E25", BLUE, "Hasil perhitungan"),
        ("G24:H25", GREEN, "Data referensi"),
        ("I24:J25", RED, "Perlu perhatian"),
    ]
    for area, color, text in legends:
        ws.merge_cells(area)
        c = ws[area.split(":")[0]]
        c.value = text
        c.fill = PatternFill("solid", fgColor=color)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    for col in range(1, 11):
        ws.column_dimensions[get_column_letter(col)].width = 15
    ws.freeze_panes = "A4"


def make_master_data(wb):
    ws = wb.create_sheet("MASTER DATA")
    title(ws, "MASTER DATA", "Daftar alat dan nilai acuan. Hanya pengelola data yang boleh mengubah sheet ini.")
    ws.sheet_view.showGridLines = False
    headers = ["Asset ID", "Nama alat", "Lokasi", "Unit", "Nilai acuan", "Toleransi +/-", "Aktif", "Penanggung jawab"]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    style_header(ws[4][:8])
    assets = [
        ("AST-001", "Pressure Gauge 0 to 10 bar", "Workshop A", "bar", 5.00, 0.10, "Yes", "Engineering"),
        ("AST-002", "Temperature Sensor", "Line 1", "deg C", 100.00, 0.50, "Yes", "Engineering"),
        ("AST-003", "Flow Meter", "Utility Area", "L/min", 50.00, 1.00, "Yes", "Maintenance"),
        ("AST-004", "Digital Caliper", "QC Laboratory", "mm", 25.00, 0.05, "Yes", "Quality"),
        ("AST-005", "Torque Wrench", "Assembly", "N m", 80.00, 2.00, "Yes", "Quality"),
        ("AST-006", "Weighing Scale", "Warehouse", "kg", 10.00, 0.02, "Yes", "Operations"),
        ("AST-007", "Humidity Sensor", "Storage Room", "%RH", 60.00, 2.00, "Yes", "Engineering"),
        ("AST-008", "Voltage Meter", "Electrical Shop", "V", 220.00, 1.00, "Yes", "Electrical"),
    ]
    for r, row in enumerate(assets, 5):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
            ws.cell(r, c).fill = PatternFill("solid", fgColor=GREEN)
    apply_grid(ws, 5, 12, 1, 8)

    ws["J4"] = "Teknisi"
    style_header([ws["J4"]])
    technicians = ["Ayu Pratama", "Bima Santoso", "Citra Lestari", "Danu Wijaya"]
    for r, name in enumerate(technicians, 5):
        ws.cell(r, 10, name)
        ws.cell(r, 10).fill = PatternFill("solid", fgColor=GREEN)
        ws.cell(r, 10).border = border

    widths = [14, 29, 18, 12, 14, 14, 11, 18, 3, 22]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:H12"
    ws.protection.sheet = False
    ws.protection.formatCells = False
    ws.protection.sort = False
    ws.protection.autoFilter = False
    for row in ws.iter_rows(min_row=5, max_row=12, min_col=1, max_col=8):
        for cell in row:
            cell.protection = Protection(locked=False)
    for row in ws.iter_rows(min_row=5, max_row=8, min_col=10, max_col=10):
        for cell in row:
            cell.protection = Protection(locked=False)

    wb.defined_names.add(DefinedName("AssetIDs", attr_text="'MASTER DATA'!$A$5:$A$12"))
    wb.defined_names.add(DefinedName("Technicians", attr_text="'MASTER DATA'!$J$5:$J$8"))


def make_input_log(wb):
    ws = wb.create_sheet("INPUT LOG")
    title(ws, "CATATAN PEMERIKSAAN TEKNISI", "Isi sel berwarna kuning. Sel biru berisi rumus. Proteksi sengaja dimatikan pada file demo.")
    ws.sheet_view.showGridLines = False
    headers = [
        "Work Order", "Tanggal", "Teknisi", "Asset ID", "Nama alat", "Lokasi", "Unit",
        "Nilai acuan", "Toleransi +/-", "Pengukuran 1", "Pengukuran 2", "Pengukuran 3", "Rata-rata", "Selisih", "Status", "Catatan"
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    style_header(ws[4])
    ws.row_dimensions[4].height = 34

    samples = [
        ("WO-2401", date(2026, 9, 21), "Ayu Pratama", "AST-001", 5.02, 5.01, 5.03, "Pemeriksaan rutin"),
        ("WO-2402", date(2026, 9, 22), "Bima Santoso", "AST-002", 100.20, 100.10, 100.30, "Pemeriksaan rutin"),
        ("WO-2403", date(2026, 9, 23), "Citra Lestari", "AST-003", 51.20, 51.10, 51.00, "Periksa penyimpangan"),
        ("WO-2404", date(2026, 9, 24), "Danu Wijaya", "AST-004", 25.01, 25.02, 25.01, "Pemeriksaan rutin"),
        ("WO-2405", date(2026, 9, 25), "Ayu Pratama", "AST-006", 10.03, 10.04, 10.03, "Perlu diperiksa ulang"),
    ]
    input_cols = [1, 2, 3, 4, 10, 11, 12, 16]
    formula_cols = [5, 6, 7, 8, 9, 13, 14, 15]
    for r in range(5, 55):
        for c in input_cols:
            ws.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)
            ws.cell(r, c).protection = Protection(locked=False)
        for c in formula_cols:
            ws.cell(r, c).fill = PatternFill("solid", fgColor=BLUE)
            ws.cell(r, c).protection = Protection(locked=True)
        ws.cell(r, 5, f'=IF($D{r}="","",IFERROR(INDEX(\'MASTER DATA\'!$B$5:$B$12,MATCH($D{r},\'MASTER DATA\'!$A$5:$A$12,0)),"NOT FOUND"))')
        ws.cell(r, 6, f'=IF($D{r}="","",IFERROR(INDEX(\'MASTER DATA\'!$C$5:$C$12,MATCH($D{r},\'MASTER DATA\'!$A$5:$A$12,0)),"NOT FOUND"))')
        ws.cell(r, 7, f'=IF($D{r}="","",IFERROR(INDEX(\'MASTER DATA\'!$D$5:$D$12,MATCH($D{r},\'MASTER DATA\'!$A$5:$A$12,0)),""))')
        ws.cell(r, 8, f'=IF($D{r}="","",IFERROR(INDEX(\'MASTER DATA\'!$E$5:$E$12,MATCH($D{r},\'MASTER DATA\'!$A$5:$A$12,0)),""))')
        ws.cell(r, 9, f'=IF($D{r}="","",IFERROR(INDEX(\'MASTER DATA\'!$F$5:$F$12,MATCH($D{r},\'MASTER DATA\'!$A$5:$A$12,0)),""))')
        ws.cell(r, 13, f'=IF(COUNT($J{r}:$L{r})<3,"",AVERAGE($J{r}:$L{r}))')
        ws.cell(r, 14, f'=IF($M{r}="","",$M{r}-$H{r})')
        ws.cell(r, 15, f'=IF(OR($A{r}="",$B{r}="",$C{r}="",$D{r}="",COUNT($J{r}:$L{r})<3),"INCOMPLETE",IF(ABS($N{r})<=$I{r},"PASS","REVIEW"))')
        ws.cell(r, 2).number_format = "dd-mmm-yyyy"
        for c in range(8, 15):
            ws.cell(r, c).number_format = "0.000"

    for idx, sample in enumerate(samples, 5):
        wo, dt, technician, asset, r1, r2, r3, notes = sample
        values = {1: wo, 2: dt, 3: technician, 4: asset, 10: r1, 11: r2, 12: r3, 16: notes}
        for col, value in values.items():
            ws.cell(idx, col, value)

    asset_dv = DataValidation(type="list", formula1="=AssetIDs", allow_blank=True)
    tech_dv = DataValidation(type="list", formula1="=Technicians", allow_blank=True)
    date_dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2035,12,31)", allow_blank=True)
    reading_dv = DataValidation(type="decimal", operator="between", formula1="-1000000", formula2="1000000", allow_blank=True)
    for dv in [asset_dv, tech_dv, date_dv, reading_dv]:
        ws.add_data_validation(dv)
    asset_dv.add("D5:D54")
    tech_dv.add("C5:C54")
    date_dv.add("B5:B54")
    reading_dv.add("J5:L54")
    asset_dv.promptTitle = "Pilih Asset ID"
    asset_dv.prompt = "Pilih kode yang tersedia di MASTER DATA."
    asset_dv.errorTitle = "Asset ID tidak ditemukan"
    asset_dv.error = "Gunakan salah satu kode dari daftar."
    asset_dv.errorStyle = "stop"
    asset_dv.showErrorMessage = True

    red_fill = PatternFill("solid", fgColor="F4CCCC")
    green_fill = PatternFill("solid", fgColor="D9EAD3")
    ws.conditional_formatting.add("O5:O54", FormulaRule(formula=['$O5="REVIEW"'], fill=red_fill))
    ws.conditional_formatting.add("O5:O54", FormulaRule(formula=['$O5="PASS"'], fill=green_fill))
    apply_grid(ws, 4, 54, 1, 16)
    widths = [16, 13, 20, 13, 27, 18, 10, 13, 13, 12, 12, 12, 13, 13, 14, 24]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = "A4:P54"
    ws.protection.sheet = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = True
    ws.protection.autoFilter = False


def make_calculation(wb):
    ws = wb.create_sheet("CALCULATION")
    title(ws, "RINCIAN PERHITUNGAN", "Sheet ini hanya untuk dibaca. Setiap baris mengikuti data pada INPUT LOG.")
    ws.sheet_view.showGridLines = False
    headers = ["Baris", "Work Order", "Asset ID", "Nilai acuan", "Toleransi", "Rata-rata", "Selisih", "Aturan", "Status"]
    for c, value in enumerate(headers, 1):
        ws.cell(4, c, value)
    style_header(ws[4][:9])
    for r in range(5, 55):
        src = r
        ws.cell(r, 1, src)
        ws.cell(r, 2, f'=\'INPUT LOG\'!A{src}')
        ws.cell(r, 3, f'=\'INPUT LOG\'!D{src}')
        ws.cell(r, 4, f'=\'INPUT LOG\'!H{src}')
        ws.cell(r, 5, f'=\'INPUT LOG\'!I{src}')
        ws.cell(r, 6, f'=\'INPUT LOG\'!M{src}')
        ws.cell(r, 7, f'=\'INPUT LOG\'!N{src}')
        ws.cell(r, 8, "ABS(Selisih) <= Toleransi")
        ws.cell(r, 9, f'=\'INPUT LOG\'!O{src}')
        for c in range(1, 10):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=BLUE)
        for c in range(4, 8):
            ws.cell(r, c).number_format = "0.000"
    apply_grid(ws, 4, 54, 1, 9)
    widths = [8, 16, 13, 13, 13, 13, 13, 29, 15]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:I54"
    ws.protection.sheet = False


def make_report(wb):
    ws = wb.create_sheet("REPORT")
    title(ws, "RINGKASAN HASIL", "Angka pada sheet ini dibaca langsung dari INPUT LOG.")
    ws.sheet_view.showGridLines = False
    cards = [
        ("A4:B4", "A5:B6", "Selesai diisi", '=COUNTIF(\'INPUT LOG\'!$O$5:$O$54,"PASS")+COUNTIF(\'INPUT LOG\'!$O$5:$O$54,"REVIEW")', BLUE),
        ("D4:E4", "D5:E6", "PASS", '=COUNTIF(\'INPUT LOG\'!$O$5:$O$54,"PASS")', GREEN),
        ("G4:H4", "G5:H6", "Perlu diperiksa", '=COUNTIF(\'INPUT LOG\'!$O$5:$O$54,"REVIEW")', RED),
        ("J4:K4", "J5:K6", "Belum lengkap", '=COUNTIF(\'INPUT LOG\'!$O$5:$O$54,"INCOMPLETE")-COUNTBLANK(\'INPUT LOG\'!$A$5:$A$54)', YELLOW),
    ]
    for label_area, value_area, label, formula, color in cards:
        label_start = label_area.split(":")[0]
        value_start = value_area.split(":")[0]
        ws.merge_cells(label_area)
        ws.merge_cells(value_area)
        cell = ws[label_start]
        cell.value = label
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(size=12, bold=True, color=DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=Side(style="medium", color=TEAL), right=Side(style="medium", color=TEAL), top=Side(style="medium", color=TEAL), bottom=Side(style="medium", color=TEAL))
        value_cell = ws[value_start]
        value_cell.value = formula
        value_cell.fill = PatternFill("solid", fgColor=color)
        value_cell.font = Font(size=20, bold=True, color=NAVY)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.border = Border(left=Side(style="medium", color=TEAL), right=Side(style="medium", color=TEAL), top=Side(style="medium", color=TEAL), bottom=Side(style="medium", color=TEAL))

    ws["A9"] = "Daftar pekerjaan"
    section(ws["A9"], ws["A9"].value)
    headers = ["Work Order", "Tanggal", "Teknisi", "Asset ID", "Nama alat", "Rata-rata", "Selisih", "Status"]
    for c, value in enumerate(headers, 1):
        ws.cell(10, c, value)
    style_header(ws[10][:8])
    for r in range(11, 61):
        src = r - 6
        refs = ["A", "B", "C", "D", "E", "M", "N", "O"]
        for c, ref in enumerate(refs, 1):
            ws.cell(r, c, f'=IF(\'INPUT LOG\'!$A${src}="","",\'INPUT LOG\'!${ref}${src})')
            ws.cell(r, c).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(r, 2).number_format = "dd-mmm-yyyy"
        ws.cell(r, 6).number_format = "0.000"
        ws.cell(r, 7).number_format = "0.000"
    apply_grid(ws, 10, 60, 1, 8)
    widths = [17, 14, 20, 14, 29, 13, 13, 15, 3, 16, 16]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A11"
    ws.auto_filter.ref = "A10:H60"
    ws.protection.sheet = False


def make_validation(wb):
    ws = wb.create_sheet("VALIDATION")
    title(ws, "PEMERIKSAAN HASIL", "Nilai acuan di kolom kuning dibandingkan dengan hasil hitung workbook.")
    ws.sheet_view.showGridLines = False
    headers = ["Test ID", "Work Order", "Rata-rata acuan", "Hasil workbook", "Status acuan", "Status workbook", "Hasil cek", "Catatan"]
    for c, value in enumerate(headers, 1):
        ws.cell(4, c, value)
    style_header(ws[4][:8])
    expected = [
        ("TC-001", "WO-2401", 5.020, "PASS", "Tekanan masih dalam toleransi"),
        ("TC-002", "WO-2402", 100.200, "PASS", "Suhu masih dalam toleransi"),
        ("TC-003", "WO-2403", 51.100, "REVIEW", "Aliran melewati toleransi"),
        ("TC-004", "WO-2404", 25.0133333333, "PASS", "Kaliper masih dalam toleransi"),
        ("TC-005", "WO-2405", 10.0333333333, "REVIEW", "Timbangan melewati toleransi"),
    ]
    for r, (test_id, wo, avg, status, note) in enumerate(expected, 5):
        ws.cell(r, 1, test_id)
        ws.cell(r, 2, wo)
        ws.cell(r, 3, avg)
        ws.cell(r, 4, f'=IFERROR(INDEX(\'INPUT LOG\'!$M$5:$M$54,MATCH($B{r},\'INPUT LOG\'!$A$5:$A$54,0)),"")')
        ws.cell(r, 5, status)
        ws.cell(r, 6, f'=IFERROR(INDEX(\'INPUT LOG\'!$O$5:$O$54,MATCH($B{r},\'INPUT LOG\'!$A$5:$A$54,0)),"")')
        ws.cell(r, 7, f'=IF(AND(ABS($C{r}-$D{r})<0.000001,$E{r}=$F{r}),"PASS","CHECK")')
        ws.cell(r, 8, note)
        for c in [1, 2, 3, 5, 8]:
            ws.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)
            ws.cell(r, c).protection = Protection(locked=False)
        for c in [4, 6, 7]:
            ws.cell(r, c).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(r, 3).number_format = "0.000"
        ws.cell(r, 4).number_format = "0.000"
    apply_grid(ws, 4, 9, 1, 8)
    ws["A12"] = "Syarat penerimaan"
    section(ws["A12"], ws["A12"].value)
    ws.merge_cells("A13:H15")
    ws["A13"] = (
        "Semua baris di kolom Hasil cek harus berstatus PASS. Jika ada perbedaan, catat penyebab dan persetujuannya "
        "di CHANGE LOG sebelum file diberikan kepada pengguna untuk UAT."
    )
    ws["A13"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["A13"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    widths = [13, 16, 19, 18, 18, 16, 13, 31]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A5"
    ws.protection.sheet = False


def make_change_log(wb):
    ws = wb.create_sheet("CHANGE LOG")
    title(ws, "CHANGE LOG", "Isi satu baris setiap kali struktur, rumus, atau aturan workbook berubah.")
    ws.sheet_view.showGridLines = False
    headers = ["Versi", "Tanggal", "Bagian", "Perubahan", "Alasan", "Persetujuan teknis", "Diperiksa oleh", "Status"]
    for c, value in enumerate(headers, 1):
        ws.cell(4, c, value)
    style_header(ws[4][:8])
    entries = [
        ("0.1", date(2026, 9, 26), "Struktur workbook", "Memisahkan input, master data, perhitungan, laporan, dan validasi", "Membuat contoh kerja", "Tidak berlaku, data fiktif", "Pemeriksaan mandiri", "Selesai"),
        ("0.1", date(2026, 9, 26), "Kontrol input", "Menambahkan dropdown, batas nilai, warna sel, dan proteksi formula", "Mengurangi salah input", "Tidak berlaku, data fiktif", "Pemeriksaan mandiri", "Selesai"),
        ("0.1", date(2026, 9, 26), "Perhitungan", "Menetapkan aturan contoh: ABS(selisih) <= toleransi", "Menguji alur perhitungan", "Ganti dengan aturan yang disetujui klien", "Pemeriksaan mandiri", "Selesai"),
    ]
    for r, entry in enumerate(entries, 5):
        for c, value in enumerate(entry, 1):
            ws.cell(r, c, value)
            ws.cell(r, c).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(r, 2).number_format = "dd-mmm-yyyy"
    apply_grid(ws, 4, 12, 1, 8)
    for r in range(8, 13):
        for c in range(1, 9):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=YELLOW)
            ws.cell(r, c).protection = Protection(locked=False)
    widths = [11, 14, 21, 44, 27, 32, 18, 14]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A5"


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    make_start_here(wb)
    make_master_data(wb)
    make_input_log(wb)
    make_calculation(wb)
    make_report(wb)
    make_validation(wb)
    make_change_log(wb)
    wb.properties.title = "Contoh perbaikan template Excel"
    wb.properties.subject = "Contoh input teknisi, master data, perhitungan, laporan, dan pemeriksaan hasil"
    wb.properties.creator = "[Nama Anda]"
    wb.properties.description = "Contoh dibuat ulang dari pengalaman sebelumnya. Semua data bersifat fiktif."
    wb.save(OUTPUT)

    # Reopen once to ensure the generated package is structurally readable.
    check = load_workbook(OUTPUT, data_only=False)
    required = {"START HERE", "MASTER DATA", "INPUT LOG", "CALCULATION", "REPORT", "VALIDATION", "CHANGE LOG"}
    missing = required.difference(check.sheetnames)
    if missing:
        raise RuntimeError(f"Missing sheets: {sorted(missing)}")
    check.close()
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
